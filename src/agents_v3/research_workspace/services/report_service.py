"""报告管理服务 - 增强版"""

from __future__ import annotations

import json
import uuid
from datetime import datetime
from typing import Any

from loguru import logger

from src.agents_v3.research_workspace.models import Report, ReportVersion
from src.agents_v3.research_workspace.storage import get_storage


class VersionConflictError(Exception):
    """乐观锁版本冲突"""

    def __init__(self, report_id: str, expected: int, actual: int):
        self.report_id = report_id
        self.expected = expected
        self.actual = actual
        super().__init__(f"Report {report_id} version conflict: expected v{expected}, got v{actual}")


def _report_to_db(report: Report) -> dict[str, Any]:
    """Report model → DB row (report_type + metadata JSONB)"""
    rtype = report.type
    rtype_str = rtype.value if hasattr(rtype, "value") else str(rtype)
    return {
        "report_id": report.report_id,
        "project_id": report.project_id,
        "report_type": rtype_str,
        "title": report.title,
        "content": report.content,
        "metadata": {
            "status": report.status,
            "scope": report.scope,
            "paper_ids": report.paper_ids,
            "evidence_ids": report.evidence_ids,
            "graph_node_ids": report.graph_node_ids,
            "section_sources": report.section_sources,
            "validation_result": report.validation_result,
            "exported_formats": report.exported_formats,
            "version": report.version,
        },
        "created_at": report.created_at,
        "updated_at": report.updated_at,
    }


def _db_to_report(item: dict[str, Any]) -> Report:
    """DB row → Report model"""
    meta = item.get("metadata", {}) or {}
    from src.agents_v3.research_workspace.models.enums import ReportType
    raw_type = item.get("report_type", "")
    try:
        rtype = ReportType(raw_type) if raw_type else ReportType.LITERATURE_REVIEW
    except ValueError:
        rtype = ReportType.LITERATURE_REVIEW
    return Report(
        report_id=item["report_id"],
        project_id=item.get("project_id", ""),
        type=rtype,
        title=item.get("title", ""),
        content=item.get("content", ""),
        scope=meta.get("scope", {}),
        paper_ids=meta.get("paper_ids", []),
        evidence_ids=meta.get("evidence_ids", []),
        graph_node_ids=meta.get("graph_node_ids", []),
        status=meta.get("status", "draft"),
        section_sources=meta.get("section_sources", {}),
        validation_result=meta.get("validation_result", {}),
        exported_formats=meta.get("exported_formats", []),
        metadata=item.get("metadata", {}),
        created_at=item.get("created_at", ""),
        updated_at=item.get("updated_at", ""),
        version=meta.get("version", 1),
    )


class ReportService:
    """报告 CRUD、版本管理、来源索引、校验和导出"""

    def __init__(self, storage=None):
        self.storage = storage or get_storage()

    # ── 保存与查询 ──────────────────────────────

    def save_report(self, report: Report, expected_version: int | None = None) -> Report:
        """保存报告（自动设置 section_sources 和 validation_result）

        Args:
            expected_version: 乐观锁 — 若指定，检查当前版本是否匹配，不匹配则抛 VersionConflictError
        """
        if expected_version is not None:
            existing = self.get_report(report.report_id)
            if existing and existing.version != expected_version:
                raise VersionConflictError(
                    report.report_id, expected_version, existing.version,
                )

        if not report.section_sources and report.scope.get("section_sources"):
            report.section_sources = report.scope["section_sources"]
        if not report.validation_result and report.scope.get("validation_result"):
            report.validation_result = report.scope["validation_result"]

        report.updated_at = datetime.now().isoformat()
        self.storage.upsert_item("reports", report.report_id, _report_to_db(report))
        logger.info(f"Saved report: {report.report_id} (v{report.version}, status={report.status})")
        return report

    def list_reports(
        self, project_id: str, report_type: str | None = None,
        status: str | None = None,
    ) -> list[Report]:
        query: dict[str, Any] = {"project_id": project_id}
        if report_type:
            query["report_type"] = report_type
        items = self.storage.query("reports", query)
        reports = [_db_to_report(i) for i in items]
        if status:
            reports = [r for r in reports if r.status == status]
        return reports

    def get_report(self, report_id: str) -> Report | None:
        item = self.storage.get_item("reports", report_id)
        if item:
            return _db_to_report(item)
        return None

    def update_report_status(self, report_id: str, status: str) -> Report | None:
        report = self.get_report(report_id)
        if not report:
            return None
        report.status = status
        report.updated_at = datetime.now().isoformat()
        self.storage.upsert_item("reports", report_id, _report_to_db(report))
        logger.info(f"Report {report_id} status -> {status}")
        return report

    def archive_report(self, report_id: str) -> Report | None:
        return self.update_report_status(report_id, "archived")

    def delete_report(self, report_id: str) -> bool:
        report = self.get_report(report_id)
        if not report:
            return False
        # Delete versions
        versions = self.storage.query("report_versions", {"report_id": report_id})
        for v in versions:
            self.storage.delete_item("report_versions", v["version_id"])
        # Delete report
        self.storage.delete_item("reports", report_id)
        logger.info(f"Deleted report {report_id} and {len(versions)} versions")
        return True

    # ── 版本管理 ──────────────────────────────────

    def create_version(
        self, report_id: str, content: str, reason: str
    ) -> ReportVersion | None:
        report = self.get_report(report_id)
        if not report:
            return None

        version = ReportVersion(
            version_id=f"ver_{uuid.uuid4().hex[:8]}",
            report_id=report_id,
            version_number=report.version,
            content=content,
            reason=reason,
            scope_snapshot=report.scope,
            source_snapshot={
                "paper_ids": report.paper_ids,
                "evidence_ids": report.evidence_ids,
                "graph_node_ids": report.graph_node_ids,
                "section_sources": report.section_sources,
            },
            paper_ids=report.paper_ids,
            evidence_ids=report.evidence_ids,
            validation_result=report.validation_result,
        )

        self.storage.upsert_item(
            "report_versions", version.version_id, version.model_dump()
        )

        report.content = content
        report.version += 1
        report.updated_at = datetime.now().isoformat()
        self.storage.upsert_item("reports", report_id, _report_to_db(report))

        logger.info(f"Created version {version.version_number} for report {report_id}")
        return version

    def list_versions(self, report_id: str) -> list[ReportVersion]:
        items = self.storage.query("report_versions", {"report_id": report_id})
        return sorted(
            [ReportVersion(**i) for i in items],
            key=lambda v: v.version_number,
        )

    def get_version(self, version_id: str) -> ReportVersion | None:
        item = self.storage.get_item("report_versions", version_id)
        if item:
            return ReportVersion(**item)
        return None

    def restore_version(self, report_id: str, version_id: str) -> Report | None:
        version = self.get_version(version_id)
        if not version or version.report_id != report_id:
            return None
        report = self.get_report(report_id)
        if not report:
            return None

        # Save current as version first
        self.create_version(report_id, report.content, "auto-save before restore")

        # Restore
        report.content = version.content
        report.paper_ids = version.paper_ids
        report.evidence_ids = version.evidence_ids
        report.section_sources = version.source_snapshot.get("section_sources", {})
        report.updated_at = datetime.now().isoformat()
        self.storage.upsert_item("reports", report_id, _report_to_db(report))

        logger.info(f"Restored report {report_id} to version {version.version_number}")
        return report

    def diff_versions(
        self, report_id: str, version_a: int, version_b: int
    ) -> dict[str, Any] | None:
        """比较两个版本的差异（段落级别）

        Args:
            report_id: 报告 ID
            version_a: 旧版本号
            version_b: 新版本号

        Returns:
            diff 结果字段：unified_diff, stats, paragraph_diffs
        """
        import difflib

        versions = self.list_versions(report_id)
        ver_map = {v.version_number: v for v in versions}

        # 也支持当前版本（未快照的最新内容）
        report = self.get_report(report_id)
        current_ver = report.version - 1 if report else 0
        if report and current_ver > 0 and current_ver not in ver_map:
            ver_map[current_ver] = ReportVersion(
                version_id="current",
                report_id=report_id,
                version_number=current_ver,
                content=report.content,
                reason="当前版本（未快照）",
            )

        va = ver_map.get(version_a)
        vb = ver_map.get(version_b)
        if not va or not vb:
            return None

        paras_a = va.content.split("\n\n")
        paras_b = vb.content.split("\n\n")

        # Unified diff
        unified = list(difflib.unified_diff(
            paras_a, paras_b,
            fromfile=f"v{version_a}",
            tofile=f"v{version_b}",
            lineterm="",
        ))

        # 段落级变更统计
        sm = difflib.SequenceMatcher(None, paras_a, paras_b)
        added = 0
        removed = 0
        changed = 0
        paragraph_diffs: list[dict] = []

        for tag, i1, i2, j1, j2 in sm.get_opcodes():
            if tag == "insert":
                added += i2 - i1 if False else j2 - j1
                for j in range(j1, j2):
                    paragraph_diffs.append({
                        "type": "added",
                        "paragraph": paras_b[j][:200],
                        "position": j,
                    })
            elif tag == "delete":
                removed += i2 - i1
                for i in range(i1, i2):
                    paragraph_diffs.append({
                        "type": "removed",
                        "paragraph": paras_a[i][:200],
                        "position": i,
                    })
            elif tag == "replace":
                changed += max(i2 - i1, j2 - j1)
                for i in range(i1, i2):
                    paragraph_diffs.append({
                        "type": "changed",
                        "old": paras_a[i][:200],
                        "new": paras_b[j1 + (i - i1)][:200] if j1 + (i - i1) < j2 else "",
                        "position": i,
                    })

        return {
            "report_id": report_id,
            "version_a": version_a,
            "version_b": version_b,
            "unified_diff": "\n".join(unified),
            "stats": {
                "added": added,
                "removed": removed,
                "changed": changed,
                "total_paragraphs_a": len(paras_a),
                "total_paragraphs_b": len(paras_b),
            },
            "paragraph_diffs": paragraph_diffs[:50],  # 上限 50 条
        }

    # ── 来源索引 ──────────────────────────────────

    def build_source_index(self, report_id: str) -> dict[str, Any]:
        """构建报告来源索引"""
        report = self.get_report(report_id)
        if not report:
            return {}

        # Paper sources
        paper_sources = []
        for pid in report.paper_ids:
            p = self.storage.get_item("papers", pid)
            if p:
                # Find which sections use this paper
                used_in = []
                for sid, sdata in report.section_sources.items():
                    if pid in sdata.get("paper_ids", []):
                        used_in.append(sid)
                identifiers = p.get("identifiers", {})
                dates = p.get("dates", {})
                source = p.get("source", {})
                authors_raw = p.get("authors", [])
                author_names = [a.get("name", "") if isinstance(a, dict) else str(a) for a in authors_raw]
                paper_meta = {
                    "paper_id": pid,
                    "title": p.get("title", ""),
                    "authors": author_names,
                    "year": dates.get("year"),
                    "venue": source.get("venue", ""),
                    "doi": identifiers.get("doi", ""),
                }
                from src.agents_v3.research_workspace.services.citation_formatter import CitationFormatter
                citations = {style: CitationFormatter.format(paper_meta, style) for style in ("simple", "apa", "gbt7714", "bibtex")}
                paper_sources.append({
                    **paper_meta,
                    "url": p.get("url", ""),
                    "citation_key": pid,
                    "used_in_sections": used_in,
                    "citations": citations,
                })

        # Evidence sources
        evidence_sources = []
        for eid in report.evidence_ids:
            e = self.storage.get_item("evidence_records", eid)
            if e:
                used_in = []
                for sid, sdata in report.section_sources.items():
                    if eid in sdata.get("evidence_ids", []):
                        used_in.append(sid)
                evidence_sources.append({
                    "evidence_id": eid,
                    "paper_id": e.get("paper_id", ""),
                    "topic": e.get("topic", ""),
                    "finding": e.get("finding", "")[:100] if e.get("finding") else "",
                    "limitation": e.get("limitation", "")[:100] if e.get("limitation") else "",
                    "source_quote": (e.get("source_quote") or "")[:200],
                    "used_in_sections": used_in,
                })

        return {
            "report_id": report_id,
            "project_id": report.project_id,
            "paper_sources": paper_sources,
            "evidence_sources": evidence_sources,
            "section_sources": report.section_sources,
            "created_at": datetime.now().isoformat(),
        }

    # ── 校验 ──────────────────────────────────────

    def validate_report_traceability(self, report_id: str) -> dict[str, Any]:
        """校验报告来源可追溯性"""
        report = self.get_report(report_id)
        if not report:
            return {"valid": False, "error": "report_not_found"}

        issues = []
        warnings = []

        # Check paper_ids exist
        missing_papers = []
        for pid in report.paper_ids:
            if not self.storage.get_item("papers", pid):
                missing_papers.append(pid)
        if missing_papers:
            issues.append({"code": "missing_paper", "severity": "error", "ids": missing_papers})

        # Check evidence_ids exist
        missing_evidence = []
        for eid in report.evidence_ids:
            if not self.storage.get_item("evidence_records", eid):
                missing_evidence.append(eid)
        if missing_evidence:
            issues.append({"code": "missing_evidence", "severity": "error", "ids": missing_evidence})

        # Scope violation check
        scope_papers = set(report.scope.get("paper_ids", []))
        scope_evidence = set(report.scope.get("evidence_ids", []))
        scope_violations = []
        if scope_papers:
            for pid in report.paper_ids:
                if pid not in scope_papers:
                    scope_violations.append({"type": "paper_out_of_scope", "id": pid})
        if scope_evidence:
            for eid in report.evidence_ids:
                if eid not in scope_evidence:
                    scope_violations.append({"type": "evidence_out_of_scope", "id": eid})

        # Section source coverage
        sections = report.section_sources
        if sections:
            sections_with_evidence = sum(1 for s in sections.values() if s.get("evidence_ids"))
            section_coverage = sections_with_evidence / len(sections) if sections else 0
        else:
            section_coverage = 0
            warnings.append({"code": "missing_section_sources", "message": "Report has no section_sources"})

        # Source coverage: what % of paper_ids appear in section_sources
        all_section_papers = set()
        for sdata in sections.values():
            all_section_papers.update(sdata.get("paper_ids", []))
        source_coverage = len(all_section_papers & set(report.paper_ids)) / len(report.paper_ids) if report.paper_ids else 0

        return {
            "valid": len(issues) == 0 and len(scope_violations) == 0,
            "report_id": report_id,
            "issues": issues,
            "warnings": warnings,
            "missing_paper_ids": missing_papers,
            "missing_evidence_ids": missing_evidence,
            "scope_violations": scope_violations,
            "source_coverage": source_coverage,
            "section_source_coverage": section_coverage,
            "checked_at": datetime.now().isoformat(),
        }

    # ── 导出 ──────────────────────────────────────

    def export_markdown(
        self, report_id: str,
        include_source_index: bool = True,
        include_validation: bool = False,
    ) -> str:
        report = self.get_report(report_id)
        if not report:
            return ""

        parts = []

        # Front matter
        parts.append(report.content)

        # Source index
        if include_source_index:
            source_index = self.build_source_index(report_id)
            if source_index.get("paper_sources"):
                parts.append("\n\n## 来源索引\n")
                parts.append("### 使用论文\n")
                for p in source_index["paper_sources"]:
                    authors = ", ".join(p.get("authors", [])[:3])
                    year = p.get("year", "")
                    title = p.get("title", "N/A")
                    doi = p.get("doi", "")
                    ref = f"- [{p['paper_id']}] {authors} ({year}). {title}."
                    if doi:
                        ref += f" DOI: {doi}"
                    parts.append(ref)

                if source_index.get("evidence_sources"):
                    parts.append("\n### 使用证据\n")
                    for e in source_index["evidence_sources"]:
                        line = f"- [{e['evidence_id']}] 论文 {e['paper_id']}"
                        if e.get("finding"):
                            line += f" | 发现: {e['finding'][:60]}"
                        if e.get("source_quote"):
                            line += f" | 引用: {e['source_quote'][:60]}"
                        parts.append(line)

        # Metadata footer
        parts.append(f"\n\n---\n")
        parts.append(f"**报告ID**: {report.report_id}")
        parts.append(f"**项目**: {report.project_id}")
        parts.append(f"**类型**: {report.type.value}")
        parts.append(f"**状态**: {report.status}")
        parts.append(f"**版本**: {report.version}")
        parts.append(f"**论文数**: {len(report.paper_ids)}")
        parts.append(f"**证据数**: {len(report.evidence_ids)}")
        parts.append(f"**生成时间**: {report.created_at}")

        # Validation
        if include_validation and report.validation_result:
            vr = report.validation_result
            parts.append(f"\n**质量校验**: {'通过' if vr.get('valid', True) else '未通过'}")
            if vr.get("source_coverage") is not None:
                parts.append(f"**来源覆盖率**: {vr['source_coverage']:.0%}")

        result = "\n".join(parts)

        # Track export
        if "markdown" not in report.exported_formats:
            report.exported_formats.append("markdown")
            self.storage.upsert_item("reports", report_id, _report_to_db(report))

        return result

    def export_json(self, report_id: str) -> str:
        """导出完整报告 JSON"""
        report = self.get_report(report_id)
        if not report:
            return "{}"

        source_index = self.build_source_index(report_id)

        export_data = {
            "report_id": report.report_id,
            "project_id": report.project_id,
            "type": report.type.value,
            "title": report.title,
            "status": report.status,
            "version": report.version,
            "content": report.content,
            "scope": report.scope,
            "paper_ids": report.paper_ids,
            "evidence_ids": report.evidence_ids,
            "graph_node_ids": report.graph_node_ids,
            "section_sources": report.section_sources,
            "validation_result": report.validation_result,
            "source_index": source_index,
            "metadata": report.metadata,
            "created_at": report.created_at,
            "updated_at": report.updated_at,
        }

        if "json" not in report.exported_formats:
            report.exported_formats.append("json")
            self.storage.upsert_item("reports", report_id, _report_to_db(report))

        return json.dumps(export_data, ensure_ascii=False, indent=2)

    def export_docx(self, report_id: str, include_source_index: bool = True) -> bytes | None:
        """导出报告为 DOCX 格式

        使用 python-docx 将 Markdown 内容转换为结构化 Word 文档。
        保留标题层级、段落、列表和简单表格。
        """
        report = self.get_report(report_id)
        if not report:
            return None

        try:
            from docx import Document
            from docx.shared import Pt, Inches
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            logger.warning("python-docx not installed, DOCX export unavailable")
            return None

        doc = Document()

        # 标题
        doc.add_heading(report.title or report.type.value, level=0)
        doc.add_paragraph(f"项目: {report.project_id} | 版本: {report.version} | 状态: {report.status}")

        # 内容：逐行解析 Markdown
        lines = report.content.split("\n")
        i = 0
        while i < len(lines):
            line = lines[i].rstrip()

            if line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("- ") or line.startswith("* "):
                doc.add_paragraph(line[2:], style="List Bullet")
            elif line.startswith("  - ") or line.startswith("  * "):
                doc.add_paragraph(line[4:], style="List Bullet 2")
            elif line.startswith("|") and "|" in line[1:]:
                # 简单表格检测：收集连续的 | 行
                table_lines = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    table_lines.append(lines[i].strip())
                    i += 1
                i -= 1  # 回退一行，外层循环会 +1
                if len(table_lines) >= 2:
                    # 解析表头和数据行
                    header = [c.strip() for c in table_lines[0].split("|")[1:-1]]
                    # 跳过分隔行 (|---|---|)
                    data_rows = []
                    for tl in table_lines[2:]:
                        cells = [c.strip() for c in tl.split("|")[1:-1]]
                        if cells:
                            data_rows.append(cells)
                    if header and data_rows:
                        table = doc.add_table(rows=1 + len(data_rows), cols=len(header))
                        table.style = "Table Grid"
                        for j, h in enumerate(header):
                            table.rows[0].cells[j].text = h
                        for ri, row in enumerate(data_rows):
                            for ci, cell in enumerate(row):
                                if ci < len(header):
                                    table.rows[ri + 1].cells[ci].text = cell
            elif line.startswith("---"):
                doc.add_paragraph("─" * 40)
            elif line.strip():
                doc.add_paragraph(line)

            i += 1

        # 来源索引
        if include_source_index:
            source_index = self.build_source_index(report_id)
            if source_index.get("paper_sources"):
                doc.add_heading("来源索引", level=1)
                doc.add_heading("使用论文", level=2)
                for p in source_index["paper_sources"]:
                    authors = ", ".join(p.get("authors", [])[:3])
                    year = p.get("year", "")
                    title = p.get("title", "N/A")
                    doi = p.get("doi", "")
                    ref = f"[{p['paper_id']}] {authors} ({year}). {title}."
                    if doi:
                        ref += f" DOI: {doi}"
                    doc.add_paragraph(ref, style="List Bullet")

                if source_index.get("evidence_sources"):
                    doc.add_heading("使用证据", level=2)
                    for e in source_index["evidence_sources"]:
                        line = f"[{e['evidence_id']}] 论文 {e['paper_id']}"
                        if e.get("finding"):
                            line += f" | 发现: {e['finding'][:60]}"
                        doc.add_paragraph(line, style="List Bullet")

        # 写入内存
        from io import BytesIO
        buf = BytesIO()
        doc.save(buf)
        buf.seek(0)

        if "docx" not in report.exported_formats:
            report.exported_formats.append("docx")
            self.storage.upsert_item("reports", report_id, _report_to_db(report))

        return buf.getvalue()
